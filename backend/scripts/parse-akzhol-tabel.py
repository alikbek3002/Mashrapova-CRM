#!/usr/bin/env python3
"""Парсер листа «для загрузки» (табель Акжол/дзюдо, июль 2026) → JSON."""
import json, re, sys, datetime
import openpyxl

SRC = "/Users/alikbekmukanbetov/Developer/Uniqum-Sport/База.xlsx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "/dev/stdout"

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb.active

rows = [[c.value for c in r] for r in ws.iter_rows(min_row=1, max_row=ws.max_row)]

def is_block_label(row):
    v = row[2]
    return (isinstance(v, str) and ":" in v
            and any(d in v.upper() for d in ["ПН", "ВТ", "СР", "ЧТ", "ПТ", "СБ"]))

def parse_date_ddmm(s, default_year=2026):
    m = re.search(r"(\d{1,2})\.(\d{1,2})(?:\.(\d{2,4}))?", s)
    if not m:
        return None
    d, mo = int(m.group(1)), int(m.group(2))
    y = m.group(3)
    y = int(y) + 2000 if y and len(y) == 2 else (int(y) if y else default_year)
    try:
        return datetime.date(y, mo, d).isoformat()
    except ValueError:
        return None

PHONE_RE = re.compile(r"\+?\s?(996[\s\d]{9,12})")

def parse_contact(raw):
    """'996777040676 Алымбек' → phone+name; 'Аниса' → name-only; 'amoCRM: ... +996..' → phone-only."""
    if not isinstance(raw, str) or not raw.strip():
        return None
    s = raw.replace("\t", " ").strip()
    m = PHONE_RE.search(s.replace(" ", "").replace("\t", ""))
    phone = ("+" + re.sub(r"\D", "", m.group(1))) if m else None
    if s.lower().startswith("amocrm"):
        return {"phone": phone, "name": None, "note": s} if phone else None
    name = PHONE_RE.sub("", s)
    name = re.sub(r"[\d+]+", "", name).strip(" ,. ")
    name = re.sub(r"\s+", " ", name) or None
    if not phone and not name:
        return None
    return {"phone": phone, "name": name}

blocks = []
i = 0
while i < len(rows):
    if not is_block_label(rows[i]):
        i += 1
        continue
    label = rows[i][2].strip()
    capacity = (rows[i][3] or "").strip() if isinstance(rows[i][3], str) else None
    header = rows[i + 1]
    dates = {}
    for col in range(5, 19):
        v = header[col]
        if isinstance(v, datetime.datetime):
            dates[col] = v.date().isoformat()
    kids = []
    j = i + 2
    while j < len(rows):
        row = rows[j]
        if is_block_label(row):
            break
        name = row[2]
        if isinstance(name, str) and name.strip() == "итого:":
            j += 1
            break
        if isinstance(name, str) and name.strip():
            kid = {"row": j + 1, "name": re.sub(r"\s+", " ", name).strip()}
            kid["manager"] = row[0].strip() if isinstance(row[0], str) else None
            start_text = row[3].strip() if isinstance(row[3], str) else None
            end_text = row[4].strip() if isinstance(row[4], str) else None
            kid["start_text"] = start_text
            kid["end_text"] = end_text
            kid["start_date"] = parse_date_ddmm(start_text) if start_text and start_text.startswith("с ") else None
            kid["months"] = None
            kid["lessons_pack_from_text"] = None
            if start_text:
                m = re.search(r"на\s+(\d+)\s*мес", start_text)
                if m: kid["months"] = int(m.group(1))
                m = re.search(r"на\s+(\d+)\s*трен", start_text)
                if m: kid["lessons_pack_from_text"] = int(m.group(1))
            kid["end_date"] = parse_date_ddmm(end_text) if end_text and "АБЗ" in end_text else None
            att = {}
            for col, dt in dates.items():
                v = row[col]
                if v == 1 or v == 1.0:
                    att[dt] = "present"
                elif isinstance(v, str) and "зам" in v.lower():
                    att[dt] = "excused"
            kid["attendance"] = att
            kid["total_marks"] = row[19] if isinstance(row[19], (int, float)) else None
            # контакт родителя: col22; иногда там же лежит амо-ссылка
            kid["amo_link"] = None
            kid["amo_note"] = None
            parent = None
            c22 = row[22]
            if isinstance(c22, str) and c22.strip().startswith("http"):
                kid["amo_link"] = c22.strip()
            else:
                parent = parse_contact(c22)
            # col26: ссылка | телефон-контакт | amoCRM-текст
            c26 = row[26]
            if isinstance(c26, str):
                s26 = c26.strip()
                if s26.startswith("http"):
                    kid["amo_link"] = kid["amo_link"] or s26
                else:
                    c = parse_contact(s26)
                    if c and not parent:
                        parent = c
                    elif c and c.get("note"):
                        kid["amo_note"] = c["note"]
                    elif s26.lower().startswith("amocrm"):
                        kid["amo_note"] = s26
            if parent and parent.get("note"):
                kid["amo_note"] = kid["amo_note"] or parent.pop("note")
            kid["parent"] = parent
            kid["in_whatsapp_group"] = bool(row[23]) if row[23] is not None else None
            kid["service_call"] = bool(row[20]) if row[20] is not None else None
            kid["coach_note"] = row[24].strip() if isinstance(row[24], str) and row[24].strip() else None
            fz = row[25]
            if isinstance(fz, (int, float)):
                kid["freeze_left"] = int(fz)
            elif isinstance(fz, str) and re.search(r"\d+", fz):
                kid["freeze_left"] = int(re.search(r"\d+", fz).group(0))
            else:
                kid["freeze_left"] = None
            kid["has_contract"] = bool(row[27]) if row[27] is not None else None
            kid["pay_date"] = row[28].date().isoformat() if isinstance(row[28], datetime.datetime) else None
            kid["pay_sum"] = float(row[29]) if isinstance(row[29], (int, float)) else None
            kid["package"] = int(row[30]) if isinstance(row[30], (int, float)) else None
            kid["lessons_left"] = int(row[31]) if isinstance(row[31], (int, float)) else None
            kid["manager_note"] = row[32].strip() if isinstance(row[32], str) and row[32].strip() else None
            kids.append(kid)
        j += 1
    blocks.append({"label": label, "capacity": capacity, "dates": sorted(dates.values()), "kids": kids})
    i = j

data = {"coach": "Акжол", "section": "Дзюдо", "month": "2026-07", "blocks": blocks}
with open(OUT, "w") as f:
    json.dump(data, f, ensure_ascii=False, indent=1)

def e(*a): print(*a, file=sys.stderr)
total = 0
names = {}
managers = {}
parent_names = set()
no_card = []
stats = {"amo_links": 0, "payments": 0, "phones": 0, "notes": 0}
for b in blocks:
    e(f"{b['label']:28} kids={len(b['kids']):3} dates={len(b['dates'])} att_marks={sum(len(k['attendance']) for k in b['kids'])}")
    for k in b["kids"]:
        total += 1
        key = k["name"].lower().replace("ё", "е")
        names.setdefault(key, []).append(b["label"])
        if k["manager"]: managers[k["manager"]] = managers.get(k["manager"], 0) + 1
        if k["parent"] and k["parent"].get("name"): parent_names.add(k["parent"]["name"])
        if k["amo_link"]: stats["amo_links"] += 1
        if k["pay_date"] and k["pay_sum"]: stats["payments"] += 1
        if k["parent"] and k["parent"].get("phone"): stats["phones"] += 1
        if k["coach_note"] or k["manager_note"] or k["amo_note"]: stats["notes"] += 1
        if not k["start_date"] and not k["end_date"]:
            no_card.append(f"{k['name']} (row {k['row']}, {b['label']})")
e(f"\nTOTAL kids: {total}")
e("managers:", managers)
e("stats:", stats)
dups = {n: bs for n, bs in names.items() if len(bs) > 1}
e("duplicates:", dups or "none")
e("no card dates:", no_card)
e("\nparent names:", sorted(parent_names, key=str.lower))
