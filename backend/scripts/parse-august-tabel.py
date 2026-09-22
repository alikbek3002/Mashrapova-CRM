#!/usr/bin/env python3
"""Парсер прод-табеля «Для загрузки 08.08.xlsx» (16 листов-тренеров) → JSON.

Каждый лист: блоки групп (секция → расписание+вместимость → шапка → дети).
Колонки мапятся по ТЕКСТУ заголовка каждого блока (в части листов колонки
сдвинуты). Кросс-фикс «пакет↔сумма»: деньги в «пакете» переезжают в сумму,
пакет ≤100 в «сумме» — в пакет.
"""
import json, re, sys, datetime
import openpyxl

OUT = sys.argv[1] if len(sys.argv) > 1 else "/Users/alikbekmukanbetov/Developer/Uniqum-Sport/backend/scripts/data/august-tabel-2026-08.json"
# argv[2] — исходный xlsx (по умолчанию первичный файл загрузки 08.08).
# Для обновлений передаём выгрузку Google-таблицы табелей.
SRC = sys.argv[2] if len(sys.argv) > 2 else "/Users/alikbekmukanbetov/Developer/Uniqum-Sport/Для загрузки 08.08.xlsx"

DOW = {"ПН": 1, "ВТ": 2, "СР": 3, "ЧТ": 4, "ПТ": 5, "СБ": 6, "ВС": 0}
SCHED_RE = re.compile(r"(ПН|ВТ|СР|ЧТ|ПТ|СБ|ВС)[\s.,]", re.I)
PHONE_RE = re.compile(r"996\d{9}")

warnings = []
def warn(msg):
    warnings.append(msg)

def cell(v):
    if v is None: return None
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    s = str(v).strip()
    return s if s != "" else None

def norm_hdr(s):
    return re.sub(r"\s+", " ", str(s).lower().replace(".", " ")).strip()

def parse_time(s):
    m = re.match(r"(\d{1,2})[:.](\d{2})", s.strip())
    if not m: return None
    return f"{int(m.group(1)):02d}:{m.group(2)}"

def minutes(t):
    h, m = t.split(":"); return int(h) * 60 + int(m)

def parse_schedule(label):
    """'ВТ.ЧТ 15:30 - 17:00 / СБ 9:30 - 11:00' → [[dow, 'HH:MM', dur], ...]"""
    out = []
    for part in re.split(r"/", label):
        days = [DOW[d.upper()] for d in re.findall(r"(?i)\b(ПН|ВТ|СР|ЧТ|ПТ|СБ|ВС)\b", part)]
        times = re.findall(r"(\d{1,2}[:.]\d{2})", part)
        if not days or not times: continue
        start = parse_time(times[0])
        dur = 60
        if len(times) >= 2:
            end = parse_time(times[1])
            if end and start:
                d = minutes(end) - minutes(start)
                if 30 <= d <= 180: dur = d
        for d in days:
            out.append([d, start, dur])
    return out

def parse_start(raw):
    """'с 17.07 на 3 мес' → (date, months, lessons, note)"""
    if not raw: return None, None, None, None
    s = str(raw).strip()
    m = re.search(r"(\d{1,2})\.(\d{1,2})", s)
    date = None
    if m:
        dd, mm = int(m.group(1)), int(m.group(2))
        try:
            date = datetime.date(2026, mm, dd).isoformat()
        except ValueError:
            pass
    months = None
    mm2 = re.search(r"на\s*(\d+)\s*мес", s)
    if mm2: months = int(mm2.group(1))
    lessons = None
    mt = re.search(r"за\s*(\d+)\s*тр", s) or re.search(r"\+\s*(\d+)\s*тр", s)
    if mt: lessons = int(mt.group(1))
    note = None
    if not date and not months:
        note = s  # «бронь на сентябрь», «разово ходят» и т.п.
    return date, months, lessons, note

def parse_end(raw):
    if not raw: return None
    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{1,4})", str(raw))
    if not m: return None
    dd, mm, yy = int(m.group(1)), int(m.group(2)), m.group(3)
    y = int(yy)
    if y < 100: y = 2000 + y if y > 6 else 2026  # «28.08.6» → 2026
    try:
        return datetime.date(y, mm, dd).isoformat()
    except ValueError:
        return None

def parse_contact(raw):
    if not raw: return None
    s = str(raw).replace("\t", " ").strip()
    digits = re.sub(r"\D", "", s)
    m = PHONE_RE.search(digits)
    phone = "+" + m.group(0) if m else None
    name = re.sub(r"[\d+()\-]+", " ", s)
    name = re.sub(r"\s+", " ", name).strip(" ,.") or None
    if not phone and not name: return None
    return {"phone": phone, "name": name}

def as_number(v):
    if v is None: return None
    if isinstance(v, datetime.date): return None
    try:
        f = float(str(v).replace(",", "."))
        return int(f) if f == int(f) else f
    except ValueError:
        return None

def as_date(v):
    if isinstance(v, datetime.date): return v.isoformat()
    if v is None: return None
    m = re.search(r"(\d{1,2})\.(\d{1,2})(?:\.(\d{2,4}))?", str(v))
    if not m: return None
    y = m.group(3)
    y = (2000 + int(y)) if y and len(y) == 2 else (int(y) if y else 2026)
    try:
        return datetime.date(y, int(m.group(2)), int(m.group(1))).isoformat()
    except ValueError:
        return None

wb = openpyxl.load_workbook(SRC, data_only=True)
sheets_out = []
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = [[cell(c) for c in r] for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=40, values_only=True)]
    coach = (rows[0][2] or sheet_name).strip() if rows and len(rows[0]) > 2 else sheet_name.strip()
    headers = [i for i, r in enumerate(rows)
               if r[0] and norm_hdr(r[0]) == "менеджер" and r[2] and norm_hdr(r[2]) == "фио"]
    # метаданные блока: расписание ищем вверх от шапки, секцию — над расписанием
    metas = []
    for h in headers:
        sched_row = section = cap = label_note = None
        for j in range(h - 1, max(h - 5, -1), -1):
            v = rows[j][2]
            if v and isinstance(v, str) and SCHED_RE.search(v) and ":" in v:
                sched_row = j
                cap = rows[j][3]
                for j2 in range(j - 1, max(j - 4, -1), -1):
                    v2 = rows[j2][2]
                    if v2 and isinstance(v2, str) and not SCHED_RE.search(v2):
                        section = v2.strip(); break
                break
        if sched_row is None:
            # блок без строки расписания (напр. «Мини группа»): сохраняем подпись
            # над шапкой — иначе группа обезличится и сольётся с соседней.
            for j in range(h - 1, max(h - 5, -1), -1):
                v = rows[j][2]
                if v and isinstance(v, str) and not norm_hdr(v).startswith("итого"):
                    label_note = v.strip(); break
        metas.append({"h": h, "sched_row": sched_row, "section": section, "cap": cap, "label_note": label_note})
    # секция-наследование: блок без секции берёт секцию предыдущего блока листа
    for i, m in enumerate(metas):
        if not m["section"]:
            for j in range(i - 1, -1, -1):
                if metas[j]["section"]: m["section"] = metas[j]["section"]; break
        if not m["section"]:
            warn(f"{sheet_name}: блок {i+1} без секции")

    blocks = []
    for bi, meta in enumerate(metas):
        h = meta["h"]
        end = metas[bi + 1]["sched_row"] if bi + 1 < len(metas) and metas[bi + 1]["sched_row"] else (
            metas[bi + 1]["h"] - 1 if bi + 1 < len(metas) else len(rows))
        hdr = rows[h]
        col = {}
        dates = []  # [(iso, col_idx)]
        for ci, v in enumerate(hdr):
            if v is None: continue
            if isinstance(v, datetime.date):
                dates.append((v.isoformat(), ci)); continue
            n = norm_hdr(v)
            if n == "менеджер": col["manager"] = ci
            elif n == "фио": col["fio"] = ci
            elif "начало" in n: col["start"] = ci
            elif "конец" in n: col["end"] = ci
            elif "остаток тренировок" in n: col["remainder"] = ci
            elif "остаток замороз" in n: col["freeze_left"] = ci
            elif "номер родителя" in n: col["phone"] = ci
            elif "амо" in n: col["amo"] = ci
            elif "договор" in n: col["contract"] = ci
            elif "дата оплаты" in n: col["pay_date"] = ci
            elif "сумма оплаты" in n: col["pay_sum"] = ci
            elif "пакет" in n: col["packet"] = ci
            elif "примечания от тренера" in n: col["coach_note"] = ci
            elif "примечания от менеджера" in n: col["manager_note"] = ci

        raw_label = rows[meta["sched_row"]][2] if meta["sched_row"] is not None else ""
        # только августовские колонки; блоки без августа — устаревшие шаблоны
        aug_dates = [(d, ci) for d, ci in dates if "2026-08-01" <= d <= "2026-08-31"]
        dropped_dates = [(d, ci) for d, ci in dates if (d, ci) not in aug_dates]
        if not aug_dates:
            warn(f"{sheet_name} «{raw_label or 'без лейбла'}»: блок без августовских дат ({dates[0][0] if dates else '—'}…) — пропущен")
            continue
        for d, ci in dropped_dates:
            marks = sum(1 for r in rows[h + 1:end] if ci < len(r) and r[ci] is not None and str(r[ci]).strip() in ("1", "1.0", "1?", "зам"))
            if marks:
                warn(f"{sheet_name} «{raw_label}»: колонка {d} вне августа с {marks} отметками — ОТБРОШЕНА, проверить вручную")
        dates = aug_dates
        sched = parse_schedule(raw_label or "")
        date_dows = {datetime.date.fromisoformat(d).isoweekday() % 7 for d, _ in dates}
        sched_dows = {s[0] for s in sched}
        if sched and sched_dows.isdisjoint(date_dows):
            # лейбл целиком противоречит датам (напр. «ПН.СР.ПТ», а даты СБ.ВС) — даты авторитетнее
            base_t, base_d = sched[0][1], sched[0][2]
            sched = [[dw, base_t, base_d] for dw in sorted(date_dows)]
            warn(f"{sheet_name} «{raw_label}»: дни лейбла не совпадают с датами, расписание взято из дат {sorted(date_dows)}")
        else:
            missing = date_dows - sched_dows
            if missing and sched:
                for dmiss in sorted(missing):
                    sched.append([dmiss, sched[0][1], sched[0][2]])
                warn(f"{sheet_name} «{raw_label}»: дни {sorted(missing)} не в расписании, добавлены со временем {sched[0][1]}")
        cap_n = None
        if meta["cap"]:
            nums = re.findall(r"\d+", str(meta["cap"]))
            if nums: cap_n = max(int(x) for x in nums)

        kids = []
        for r in rows[h + 1:end]:
            fio = r[col.get("fio", 2)]
            if not fio or not isinstance(fio, str): continue
            f = fio.strip()
            if not f or norm_hdr(f).startswith("итого"): continue
            if SCHED_RE.search(f) and ":" in f: continue
            get = lambda key: (r[col[key]] if key in col and col[key] < len(r) else None)
            manager = get("manager")
            att = {}
            for d, ci in dates:
                v = r[ci] if ci < len(r) else None
                if v is None: continue
                sv = str(v).strip().lower()
                if sv in ("1", "1.0", "1?"): att[d] = "present"
                elif sv == "зам": att[d] = "excused"
            has_signal = manager or get("start") or get("end") or att or get("phone") or get("pay_sum")
            if not has_signal: continue
            start_date, months, lessons_txt, start_note = parse_start(get("start"))
            packet = as_number(get("packet"))
            pay_sum = as_number(get("pay_sum"))
            pay_date = as_date(get("pay_date"))
            # кросс-фикс сдвинутых колонок: деньги в «пакете», пакет в «сумме»
            if packet is not None and packet > 100:
                if pay_sum is None: pay_sum = packet
                packet = None
            if pay_sum is not None and pay_sum <= 100:
                if packet is None: packet = int(pay_sum)
                pay_sum = None
            remainder = as_number(get("remainder"))
            if remainder is not None and remainder > 100:
                warn(f"{sheet_name}/{f}: остаток {remainder} > 100, обнулён в None")
                remainder = None
            kids.append({
                "manager": str(manager).strip().capitalize() if manager else None,
                "fio": re.sub(r"\s+", " ", f),
                "start_raw": str(get("start") or "").strip() or None,
                "start_date": start_date, "months": months,
                "lessons_from_text": lessons_txt, "start_note": start_note,
                "end_date": parse_end(get("end")),
                "attendance": att,
                "parent": parse_contact(get("phone")),
                "freeze_left": as_number(get("freeze_left")),
                "amo": (str(get("amo")).strip() if get("amo") and "http" in str(get("amo")) else None),
                "has_contract": str(get("contract") or "").strip().lower() in ("true", "1", "да", "есть"),
                "pay_date": pay_date, "pay_sum": pay_sum,
                "packet": packet, "remainder": remainder,
                "coach_note": (str(get("coach_note")).strip() or None) if get("coach_note") else None,
                "manager_note": (str(get("manager_note")).strip() or None) if get("manager_note") else None,
            })
        blocks.append({
            "section_raw": (meta["section"] or "").strip(),
            "schedule_raw": (raw_label or "").strip(),
            "label_note": meta.get("label_note"),
            "capacity": cap_n,
            "sched": sched,
            "dates": [d for d, _ in dates],
            "kids": kids,
        })
    sheets_out.append({"sheet": sheet_name.strip(), "coach": coach, "blocks": blocks})

out = {"generated_from": SRC, "cutoff": "2026-08-08", "sheets": sheets_out, "warnings": warnings}
with open(OUT, "w", encoding="utf-8") as fh:
    json.dump(out, fh, ensure_ascii=False, indent=1)

n_groups = sum(len(s["blocks"]) for s in sheets_out)
n_kids = sum(len(b["kids"]) for s in sheets_out for b in s["blocks"])
print(f"Листов: {len(sheets_out)}, групп: {n_groups}, строк детей: {n_kids}")
print(f"Предупреждений: {len(warnings)}")
for w in warnings: print("  !", w)
print(f"JSON → {OUT}")
